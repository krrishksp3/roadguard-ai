import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from doc_helpers import (
    set_cell_background,
    set_cell_margins,
    set_table_borders,
    add_callout,
    format_row,
    add_custom_heading,
    add_body_p,
    add_bullet_p,
    add_code_block,
)

def add_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = header_p.add_run("ROADGUARD AI | SIH Internal Hackathon – MB-04 | Team YuvaTech")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        # Footer
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = footer_p.add_run("ROADGUARD AI — Technical Project Documentation — Page ")
        r_ftr.font.name = "Calibri"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        # Add Page Number field
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        footer_p._p.append(fldSimple)
        
        r_of = footer_p.add_run(" of ")
        r_of.font.name = "Calibri"
        r_of.font.size = Pt(8.5)
        r_of.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        fldSimple2 = OxmlElement('w:fldSimple')
        fldSimple2.set(qn('w:instr'), 'NUMPAGES')
        footer_p._p.append(fldSimple2)

def build_cover_page(doc):
    # Space before title
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(36)
    
    # Badge
    badge_p = doc.add_paragraph()
    badge_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_badge = badge_p.add_run("SMART INDIA HACKATHON 2026 | INTERNAL HACKATHON")
    r_badge.font.name = "Calibri"
    r_badge.font.size = Pt(10)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("ROADGUARD AI")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(32)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("AI-Powered Smart Road Safety, Dynamic Risk Prioritization & Civic Accountability Platform")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Decorative Divider Table
    div_table = doc.add_table(rows=1, cols=1)
    div_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    div_cell = div_table.cell(0, 0)
    div_cell.width = Inches(3.0)
    set_cell_background(div_cell, "2563EB")
    set_cell_margins(div_cell, top=20, bottom=20, left=0, right=0)
    
    # Spacer
    sp2 = doc.add_paragraph()
    sp2.paragraph_format.space_before = Pt(36)
    
    # Metadata Block Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    rows_data = [
        ("Problem Statement Code", "SIH Internal Hackathon – MB-04"),
        ("Theme / Category", "Smart Infrastructure & Logistics"),
        ("Project Name", "ROADGUARD AI"),
        ("Team Name", "YuvaTech"),
        ("Demonstration Jurisdiction", "Meerut Municipal Region (UP PWD / Nagar Nigam / NHAI)"),
        ("Document Release & Date", "Version 1.0 (March 2026) | Full Engineering Documentation"),
    ]
    
    for i, (k, v) in enumerate(rows_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.3)
        c1.width = Inches(4.2)
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
    set_table_borders(meta_table, "CBD5E1")
    
    # Spacer
    sp3 = doc.add_paragraph()
    sp3.paragraph_format.space_before = Pt(40)
    
    # Institution footer note
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("Submitted for Smart India Hackathon Internal Evaluation & Technical Review\nTeam YuvaTech — Advanced Agentic & Civic Engineering Division")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(9)
    r_inst.font.italic = True
    r_inst.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    doc.add_page_break()

print("Cover page builder ready.")
