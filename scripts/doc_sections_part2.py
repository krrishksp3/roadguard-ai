import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from doc_helpers import (
    set_cell_background,
    set_cell_margins,
    set_table_borders,
    add_callout,
    format_row,
    add_custom_heading,
    add_body_p,
    add_bullet_p,
    add_code_block,
)

def render_section_16_to_28(doc):
    # SECTION 16: AUTHENTICATION & SECURITY
    add_custom_heading(doc, "16. Authentication, Role Security & Data Protection", level=1)
    add_body_p(doc, "Security in ROADGUARD AI is implemented through defense-in-depth across the client, network, application, and database tiers:")
    
    add_bullet_p(doc, "User passwords are encrypted prior to database persistence using bcryptjs with 10 salt rounds. Cleartext credentials are never written to logs or transmitted across API responses.", bold_prefix="Cryptographic Password Hashing: ")
    add_bullet_p(doc, "Sessions are managed via compact, stateless JSON Web Tokens (JWT) signed with HMAC-SHA256 (`HS256`). Tokens carry user identity (`id`), email, role (`CITIZEN`, `AUTHORITY`, `ADMIN`), and department affiliation, expiring automatically after 7 days.", bold_prefix="Stateless JWT Token Issuance: ")
    add_bullet_p(doc, "Sensitive administrative actions (such as status progression, after-repair evidence uploads, and verification verdicts) are guarded by `requireRoles(['AUTHORITY', 'ADMIN'])` middleware. Unauthorized access attempts yield immediate HTTP 403 Forbidden envelopes.", bold_prefix="Role-Based Route Interception: ")
    add_bullet_p(doc, "All database queries utilize Prisma Client parameterized statements, fully insulating the data layer from SQL injection vulnerabilities.", bold_prefix="SQL Injection Elimination: ")
    add_bullet_p(doc, "The `/api/upload` endpoint implements strict Multer disk-storage rules, restricting file sizes to 10MB and enforcing strict MIME filters (`image/jpeg`, `image/png`, `image/webp`, `image/svg+xml`). Arbitrary file extensions or executable binaries are rejected with explicit HTTP 400 errors.", bold_prefix="File Upload Sanitization: ")
    add_bullet_p(doc, "Zero private keys, production database credentials, or third-party API secrets are hardcoded in client source code. Production configuration is injected exclusively via environment variables (`JWT_SECRET`, `DATABASE_URL`, `CLOUDINARY_API_KEY`, etc.).", bold_prefix="Zero Secret Exposure Policy: ")

    # SECTION 17: AI / INTELLIGENCE MODULE
    add_custom_heading(doc, "17. AI / Intelligence Module & Multi-Provider Architecture", level=1)
    add_body_p(doc, "The AI module in ROADGUARD AI is engineered around a flexible Strategy Pattern interface (`AIProvider`), allowing the system to operate flawlessly in disconnected hackathon demo environments, on edge nodes, or connected to state-of-the-art multimodal large vision models:")
    
    add_bullet_p(doc, "Defines a clean contract requiring two core capabilities: `analyzeRoadDamage(input)` and `verifyRepair(input)`. All provider outputs are validated at runtime against strict Zod schemas (`AIAnalysisOutputSchema` and `AIVerificationOutputSchema`).", bold_prefix="1. Standardized AIProvider Interface: ")
    add_bullet_p(doc, "The primary out-of-the-box provider. Runs offline without external network or API key dependencies. Analyzes image metadata and linguistic descriptions using deterministic heuristic rules aligned with Indian Roads Congress (IRC:67 for signage and IRC:SP:72 for rural/urban asphalt design). Evaluates simulated image blur, darkness, road presence, damage type, and assigns technical civil repair recommendations.", bold_prefix="2. DemoAIProvider (Offline/Hackathon Mode): ")
    add_bullet_p(doc, "Integrates Google's `gemini-1.5-flash` or `gemini-1.5-pro` multimodal vision API. Formulates structured multimodal prompts with zero-shot civil engineering guidance, instructing the model to return strictly typed JSON evaluating pavement distress and collision risk.", bold_prefix="3. GeminiProvider (Cloud Multimodal API): ")
    add_bullet_p(doc, "Integrates OpenAI's `gpt-4o` vision model via OpenAI SDK JSON schema structured outputs, enforcing exact compliance with the platform's diagnostic schema.", bold_prefix="4. OpenAIProvider (Cloud Multimodal API): ")
    add_bullet_p(doc, "Client-side and server-side algorithms evaluate whether the uploaded photograph exhibits acceptable quality: detecting extreme darkness, severe blur, or missing road surfaces, warning citizens before they submit invalid imagery.", bold_prefix="5. Optical Quality Assurance Engine: ")

    add_callout(doc, "CRITICAL AI DISCLOSURE: In offline local development and hackathon demonstration modes, ROADGUARD AI operates using DemoAIProvider. It uses rigorous IRC civil engineering heuristics rather than an on-device convolutional neural network (CNN). Future integration with on-device YOLOv8/ONNX models is planned.", title="AI ARCHITECTURE & HONEST TECHNICAL DISCLOSURE", box_type="info")

    # SECTION 18: DYNAMIC ROAD RISK PRIORITIZATION
    add_custom_heading(doc, "18. Dynamic Road Risk Prioritization Engine", level=1)
    add_body_p(doc, "Conventional municipal systems treat all potholes equally, resulting in arbitrary repair schedules. ROADGUARD AI's `RoadRiskEngine` implements a multi-parameter mathematical formula (scaled 0 to 100) that calculates an objective dynamic risk score for every incident:")
    
    add_body_p(doc, "Dynamic Road Risk Mathematical Formulation:", bold_prefix=None)
    add_code_block(doc,
"""Overall Risk Score (0 - 100) =
    (Severity Score        × 0.25) +
    (Safety Risk Score     × 0.20) +
    (Road Importance Score × 0.20) +
    (Cluster Density Score × 0.15) +
    (Recurrence Score      × 0.10) +
    (SLA Urgency Score     × 0.10)

Clamped: Math.min(100, Math.max(1, Math.round(RawTotal)))""")

    add_body_p(doc, "Weighting Parameter Breakdown & Mathematical Scaling:")
    add_bullet_p(doc, "Scaled to 25 points. Pavement failure depth and width. Critical = 100% (25 pts), High = 75% (18.8 pts), Medium = 50% (12.5 pts), Low = 25% (6.3 pts).", bold_prefix="1. Pavement Severity (Weight: 25%): ")
    add_bullet_p(doc, "Scaled to 20 points. Direct collision, loss-of-control, or vehicle chassis damage hazard index (0-100) derived from AI analysis scaled by AI confidence.", bold_prefix="2. Road Safety Risk (Weight: 20%): ")
    add_bullet_p(doc, "Scaled to 20 points. Classification of the transit corridor. HIGHWAY = 100% (20 pts), ARTERIAL = 85% (17 pts), MAJOR_DISTRICT = 65% (13 pts), LOCAL = 40% (8 pts). Critical arteries with heavy passenger loads receive higher priority.", bold_prefix="3. Road Importance (Weight: 20%): ")
    add_bullet_p(doc, "Scaled to 15 points. Proximity density. Number of active complaints within a 150-meter radius (`nearbyReportsCount × 25`, clamped between 10 and 100). Highlights concentrated structural failure corridors.", bold_prefix="4. Cluster Density (Weight: 15%): ")
    add_bullet_p(doc, "Scaled to 10 points. Historical chronic failure index. If the road segment is an identified recurring hotspot, base score = 100% (10 pts); if road health is below 50, minimum 80% (8 pts); else baseline 20% (2 pts).", bold_prefix="5. Recurrence Factor (Weight: 10%): ")
    add_bullet_p(doc, "Scaled to 10 points. Ratio of elapsed hours since reporting against target SLA hours (`hoursSinceReported / slaTargetHours × 100`). Escalates automatically as deadlines approach.", bold_prefix="6. SLA Urgency (Weight: 10%): ")

    add_body_p(doc, "Prioritization Tiers & Explainable Diagnostic Output:")
    add_bullet_p(doc, "CRITICAL (Score >= 80) -> Immediate emergency dispatch required within 24 hours.")
    add_bullet_p(doc, "HIGH (Score >= 65 to 79) -> Arterial or high-severity defect requiring work order within 48 hours.")
    add_bullet_p(doc, "MEDIUM (Score >= 45 to 64) -> Secondary collector or moderate pavement defect; SLA target 120 hours (5 days).")
    add_bullet_p(doc, "LOW (Score < 45) -> Minor surface cosmetic defect on local street; SLA target 240 hours (10 days).")
    add_bullet_p(doc, "Explainability Engine: Rather than outputting a black-box number, the engine outputs an array of clear civil engineering rationale strings displayed directly on the Authority Dashboard (e.g., 'Critical arterial transit corridor with high passenger density', 'Cluster detected: 3 reports within 150m radius').")

    # SECTION 19: DUPLICATE & CLUSTER INTELLIGENCE
    add_custom_heading(doc, "19. Duplicate Detection & Incident Spatial Clustering", level=1)
    add_body_p(doc, "Duplicate complaints represent a major administrative drain on civic bodies. ROADGUARD AI addresses this through mathematical spatial deduplication in `src/services/duplicate/duplicateDetector.ts` using the Haversine great-circle distance formula:")
    
    add_code_block(doc,
"""Haversine Formula:
  a = sin²(Δφ/2) + cos(φ1) · cos(φ2) · sin²(Δλ/2)
  c = 2 · atan2(√a, √(1−a))
  Distance (meters) = R · c   (where R = 6,371,000 meters)""")

    add_body_p(doc, "Deduplication & Clustering Rules:")
    add_bullet_p(doc, "The detector queries active, unresolved complaints registered within the preceding 14 calendar days.", bold_prefix="1. Temporal Window: ")
    add_bullet_p(doc, "If a new submission is located within 80 meters of an active complaint AND shares the same damage classification (or 'other'), it is flagged with `isDuplicate = true` and linked via `duplicateOfId` to the primary report. Field engineers work on a single consolidated ticket while all citizens receive tracking updates.", bold_prefix="2. Proximity Duplicate Threshold (<= 80m): ")
    add_bullet_p(doc, "All unresolved reports within a 150-meter radius increment the cluster density counter (`nearbyReportsCount`). This elevates the report's Dynamic Road Risk Score, alerting authorities that an entire corridor is deteriorating.", bold_prefix="3. Spatial Density Clustering (<= 150m): ")

    # SECTION 20: ROAD HEALTH INTELLIGENCE
    add_custom_heading(doc, "20. Road Segment Health & Lifecycle Deterioration Tracking", level=1)
    add_body_p(doc, "Rather than managing isolated pothole complaints, ROADGUARD AI aggregates data at the road corridor level via `src/services/roadHealth/roadHealthService.ts`. Every physical road segment is assigned a dynamic Health Index (0 to 100):")
    
    add_code_block(doc,
"""Road Segment Health Score Formula:
  HealthScore = 100 - (OpenComplaints × 4) - (CriticalComplaints × 8) - (IsRecurringHotspot ? 15 : 0)
  Clamped: Math.max(15, Math.min(100, HealthScore))""")

    add_body_p(doc, "Corridor Health Metrics & Lifecycle Intelligence:")
    add_bullet_p(doc, "If a road segment records 3 or more complaints within a 6-month window, the system automatically flags `isRecurringHotspot = true`. This indicates deep structural subgrade failure rather than superficial surface wear.", bold_prefix="Recurring Hotspot Detection: ")
    add_bullet_p(doc, "Evaluates citizen reporting density per kilometer (`totalComplaints / lengthKm`). Categorized as LIMITED (< 0.8 / km), MODERATE (0.8 - 3.0 / km), or HIGH (> 3.0 / km) to highlight surveillance blindspots.", bold_prefix="Surveillance Coverage Estimation: ")
    add_bullet_p(doc, "When a road segment's health score drops below 50, municipal leadership is prompted to schedule full bituminous overlay resurfacing rather than spending funds on repetitive cold-mix emergency patches.", bold_prefix="Preventive Resurfacing Triggers: ")

    # SECTION 21: AI-ASSISTED REPAIR VERIFICATION
    add_custom_heading(doc, "21. AI-Assisted Repair Verification & Human-in-the-Loop Audit", level=1)
    add_body_p(doc, "To eradicate fraudulent 'paper closures', ROADGUARD AI enforces an evidence-based dual-photo verification workflow orchestrated by `src/services/verification/verificationService.ts`:")
    
    add_bullet_p(doc, "When field maintenance is completed, the supervising engineer must upload a genuine photograph of the repaired pavement section (`/api/reports/:id/after-photo`).", bold_prefix="1. Mandatory After-Repair Upload: ")
    add_bullet_p(doc, "The verification engine evaluates the Before and After imagery across three quantitative metrics: Location Match Confidence (0-100), Visible Surface Improvement Score (0-100), and Remaining Residual Damage Score (0-100).", bold_prefix="2. Computer Vision Inspection: ")
    add_bullet_p(doc, "Based on inspection thresholds, the system generates an algorithmic recommendation: PASS (acceptable patch and compaction), NEEDS_REINSPECTION (visible depressions or cracks remain), or INCONCLUSIVE (poor lighting or different camera angle).", bold_prefix="3. Algorithmic Recommendation: ")
    add_bullet_p(doc, "The AI recommendation does NOT unilaterally close the complaint. The Executive Engineer must review the side-by-side evidence on the dashboard and record an explicit decision: APPROVED or REJECTED_REINSPECT.", bold_prefix="4. Human-in-the-Loop Authority Verdict: ")

    add_callout(doc, "EVIDENCE INTEGRITY POLICY: In the demonstration environment, before-repair photos utilize authentic, reusable CC-licensed road photographs from Wikimedia Commons, while simulated after-repair imagery is clearly labeled as DEMO AI ANALYSIS. In production, field photographs are captured directly via device camera.", title="EVIDENCE INTEGRITY & VERIFICATION DISCLAIMER", box_type="warning")

    # SECTION 22: TENDER INTELLIGENCE
    add_custom_heading(doc, "22. Public Tender Intelligence & Contractor Accountability", level=1)
    add_body_p(doc, "The Tender Intelligence module bridges citizen complaints with public procurement accountability. In traditional governance, contractors who execute poor-quality road work are rarely held liable because municipal engineers lack immediate access to contract warranty terms.")
    
    add_bullet_p(doc, "Road segments are mapped to public civil tenders containing Official Tender Reference Numbers, Work Descriptions, Total Contract Values, Award Dates, and Defect Liability Periods (DLP).", bold_prefix="Tender-to-Asset Mapping: ")
    add_bullet_p(doc, "When a complaint is lodged on a road under an active Defect Liability Period (typically 36 months for bituminous asphalt), the system highlights an active warranty alert. The municipality is legally empowered to order the original contractor to repair the road at zero public cost.", bold_prefix="Defect Liability Period (DLP) Enforcement: ")
    add_bullet_p(doc, "Maintains a record of contractor performance across municipal divisions, highlighting recurring failures during warranty periods.", bold_prefix="Contractor Quality Metrics: ")

    add_callout(doc, "DEMONSTRATION NOTICE: Tender records in the project seed database (such as UP-PWD-MRT-2024-JLC-041 and M/s Meerut Highway Infrastructure Pvt. Ltd.) are realistic synthetic models calibrated for the SIH 2026 Internal Hackathon. They do not constitute official UP e-Procurement records.", title="DEMO / SYNTHETIC PROCUREMENT DATA NOTICE", box_type="info")

    # SECTION 23: MAP & LOCATION SYSTEM
    add_custom_heading(doc, "23. Geospatial Intelligence & Interactive Mapping System", level=1)
    add_body_p(doc, "Geospatial capabilities in ROADGUARD AI ensure rapid, pinpoint defect location:")
    add_bullet_p(doc, "Leverages W3C Geolocation API (`navigator.geolocation`) with high accuracy flags, acquiring device coordinates and horizontal accuracy margins.", bold_prefix="Hardware GPS Geolocation: ")
    add_bullet_p(doc, "Provides a reactive Leaflet/OpenStreetMap interface where citizens can drag a high-contrast pin to fine-tune damage location when GPS drift occurs inside dense urban corridors.", bold_prefix="Interactive Leaflet Canvas: ")
    add_bullet_p(doc, "Converts raw latitude and longitude coordinates into human-readable street addresses and landmarks using reverse geocoding heuristics.", bold_prefix="Reverse Geocoding: ")
    add_bullet_p(doc, "Coordinates are tested against known municipal road vectors, automatically identifying the responsible authority (UP PWD Provincial Division, Meerut Municipal Corporation, or NHAI PIU Meerut).", bold_prefix="Jurisdictional Snapping: ")

    # SECTION 24: PWA & OFFLINE FUNCTIONALITY
    add_custom_heading(doc, "24. Progressive Web Application (PWA) & Offline Capability", level=1)
    add_body_p(doc, "To guarantee reporting reliability across low-connectivity transit routes, ROADGUARD AI is engineered as a certified Progressive Web Application:")
    add_bullet_p(doc, "Equipped with a W3C Web App Manifest (`manifest.webmanifest`), responsive theme colors (`#1e3a8a`), and standalone display mode, allowing citizens to install the app on Android, iOS, or desktop without app store friction.", bold_prefix="Full Installability: ")
    add_bullet_p(doc, "Managed via `vite-plugin-pwa` and Workbox, precaching 39 static assets (HTML, JavaScript bundles, CSS stylesheets, and civic icons) to guarantee instant offline app shell boot.", bold_prefix="Workbox Service Worker Precaching: ")
    add_bullet_p(doc, "Implemented in `src/services/offlineSync.ts`. When a citizen submits a report without network access, the report payload and base64 photograph are stored in IndexedDB (`roadguard_offline_db`) under `pending_reports`.", bold_prefix="IndexedDB Transactional Offline Storage: ")
    add_bullet_p(doc, "A reactive network observer listens for connectivity recovery. Upon reconnect, the drainer sequentially uploads queued images, invokes the REST API, and clears synchronized drafts.", bold_prefix="Automatic Background Synchronization: ")

    # SECTION 25: DATA FLOW
    add_custom_heading(doc, "25. End-to-End Civic Data Flow Pipeline", level=1)
    add_body_p(doc, "The comprehensive data flow from citizen capture to final resolution audit is summarized below:")
    add_code_block(doc,
"""CITIZEN REPORTING STAGE
  1. Citizen opens PWA -> Captures camera photo -> Obtains GPS coordinates -> Selects damage type.
  2. Client-side quality inspection evaluates clarity -> Dispatches multipart request to /api/reports.
     [Offline Fallback: Saves to IndexedDB 'roadguard_offline_db' -> Auto-drains on network reconnect]

SERVER INGESTION & AI TRIAGE STAGE
  3. Multer uploads photo -> StorageService saves to Cloudinary / Supabase / Local disk -> Returns URL.
  4. Express calls AIService -> Classifies distress type, assesses severity & IRC maintenance recommendation.
  5. DuplicateDetector calculates Haversine distance -> Detects duplicates (<=80m) & clusters (<=150m).
  6. RoadRiskEngine computes 6-factor score (0-100) -> Generates explainable diagnostic rationale.
  7. JurisdictionService snaps incident to nearest RoadSegment & Department (UP PWD, Nagar Nigam, NHAI).
  8. SLAService computes target resolution hours -> Prisma saves RoadReport, AIAnalysis, Priority records.

MUNICIPAL ACTION & REPAIR STAGE
  9. Incident appears on Authority Operations Dashboard, ordered by Dynamic Risk Score.
  10. Executive Engineer reviews report dossier, assigns field maintenance contractor, advances status.
  11. Maintenance crew executes physical pavement repair (asphalt patch compaction / crack sealing).

VERIFICATION & LIFECYCLE AUDIT STAGE
  12. Supervising engineer captures and uploads 'After-Repair' photographic proof.
  13. VerificationService executes AI visual comparison -> Computes surface improvement & match confidence.
  14. Executive Engineer reviews dual-photo comparison -> Signs off 'APPROVED' or 'REJECTED_REINSPECT'.
  15. Status changes to RESOLVED -> StatusTimelineEvent logged -> RoadHealthService updates corridor score.""")

    # SECTION 26: COMPLETE USER JOURNEYS
    add_custom_heading(doc, "26. Complete User Journeys (Citizen & Municipal Authority)", level=1)
    add_body_p(doc, "A. Detailed Citizen User Journey:")
    add_bullet_p(doc, "Step 1 (Discovery): Citizen launches ROADGUARD AI PWA on mobile browser while encountering a hazardous pothole on Jail Chungi Road.")
    add_bullet_p(doc, "Step 2 (Quick Authentication): Citizen registers or uses one-click demo login (`citizen@roadguard.demo`).")
    add_bullet_p(doc, "Step 3 (Evidence Capture): Taps 'Report Road Hazard' -> Activates native mobile camera -> Captures high-resolution photo of the pavement defect.")
    add_bullet_p(doc, "Step 4 (Geocoding & Fine-Tuning): Device auto-detects GPS coordinates (28.9815° N, 77.7380° E). Citizen drags map pin slightly to align with lane position.")
    add_bullet_p(doc, "Step 5 (AI Live Preview): Selects damage category 'Pothole'. The optical checker displays 'Image Quality: 92% Excellent | Road Surface Detected'.")
    add_bullet_p(doc, "Step 6 (Submission): Clicks 'Submit Road Report'. Server assigns ticket ID `RG-MRT-2026-000101`, assigns Risk Score 82 (CRITICAL), and routes to UP PWD.")
    add_bullet_p(doc, "Step 7 (Tracking & Resolution): Citizen monitors `/my-reports`, receives status progression notifications, and views verified after-repair photo upon completion.")

    add_body_p(doc, "B. Detailed Municipal Authority User Journey:")
    add_bullet_p(doc, "Step 1 (Operations Console Access): Executive Engineer logs into `/authority` using departmental credentials (`authority@roadguard.demo`).")
    add_bullet_p(doc, "Step 2 (Operational Triaging): Reviews KPI cards: 18 Total Complaints, 7 Critical Incidents, 2 Pending AI Verifications. The priority table displays ticket `RG-MRT-2026-000101` at the top with Risk Score 82.")
    add_bullet_p(doc, "Step 3 (Dossier Review): Clicks ticket -> Inspects AI diagnosis, IRC remedial advice ('Cold mix or hot bituminous asphalt patch repair with mechanical compaction'), and active tender warranty alert under DLP.")
    add_bullet_p(doc, "Step 4 (Contractor Dispatch): Advances status to `IN_PROGRESS` and notifies contractor M/s Meerut Highway Infrastructure.")
    add_bullet_p(doc, "Step 5 (After-Repair Verification): Upon repair, uploads contractor post-work photo. AI verification engine generates a Visible Improvement Score of 89/100 and recommends PASS.")
    add_bullet_p(doc, "Step 6 (Final Sign-Off): Executive Engineer approves closure -> Ticket marks `RESOLVED` -> Segment health score improves automatically.")

    # SECTION 27: UI / PAGE DOCUMENTATION
    add_custom_heading(doc, "27. Comprehensive User Interface & Screen Specifications", level=1)
    add_body_p(doc, "ROADGUARD AI features 10 dedicated application pages designed with a sleek civic design system:")
    
    pages_data = [
        ("Page Name", "Route Path", "Primary Functional Elements & Widgets", "User Interactions & API Calls"),
        ("Landing Page", "/", "Civic Hero Banner, Real-Time Statistics Counters, Recent Complaints Carousel, 6-Step Workflow Infographic, CTA Buttons.", "Public landing view; loads live summary metrics from `/api/analytics/summary` and recent reports from `/api/reports`."),
        ("Login Page", "/login", "Role-aware login form, password visibility toggle, Demo Quick-Login buttons for Citizen and Authority roles.", "Submits credentials to `POST /api/auth/login`; stores signed JWT in localStorage and updates global `AuthContext`."),
        ("Register Page", "/register", "User registration form capturing Full Name, Email, Password, Phone Number, and Role selection.", "Submits user details to `POST /api/auth/register`; automatically authenticates upon successful account creation."),
        ("Create Report Page", "/report", "Camera/Gallery File Picker, Image Quality Preview, GPS Geolocation Badge, Draggable Leaflet Pin, Damage Type Grid, Description Field.", "Validates image quality client-side; submits multipart data to `POST /api/upload` then `POST /api/reports`; saves to IndexedDB if offline."),
        ("My Reports Page", "/my-reports", "Citizen complaint cards, Status Badges, Dynamic Risk Meters, SLA Countdown Timers, Offline Pending Sync Drawer.", "Fetches user records via `GET /api/reports/my-reports`; displays real-time synchronization status and retry controls."),
        ("Report Detail Page", "/reports/:id", "Before/After Side-by-Side Photo Viewer, Image Zoom Modal, AI Diagnostic Breakdown, Dynamic Risk 6-Factor Radar, Timeline, Authority Actions.", "Calls `GET /api/reports/:id`; allows authorities to patch status via `PATCH /api/reports/:id/status` and upload after-photos."),
        ("Authority Dashboard", "/authority", "KPI Metrics (Total, Critical, SLA Overdue, Avg Risk), Severity/Status Filters, Priority-Sorted Report Table, Direct Action Modals.", "Restricted to AUTHORITY role; calls `GET /api/reports` and `GET /api/analytics/summary`; supports fast inline status progression."),
        ("Public Map Page", "/map", "Full-Screen Interactive Leaflet Map, Color-Coded Risk Markers, Incident Details Drawer, Jurisdiction Boundary Filters.", "Fetches geo-tagged reports from `GET /api/reports`; renders interactive clusters and road segment overlays."),
        ("Tender Intelligence Page", "/tenders", "Public Works Tender Directory, DLP Warranty Badges, Contractor Transparency Cards, Road Segment Links.", "Calls `GET /api/tenders`; displays contract value, work duration, DLP status, and linked road health scores."),
        ("Road Health Page", "/road-health", "Road Segment Health Cards (0-100), Chronic Hotspot Badges, Health Distribution Chart, Maintenance Action Alerts.", "Calls `GET /api/road-health` and `/api/road-health/hotspots`; displays segment metrics and preventive maintenance triggers."),
    ]
    
    tbl_pg = doc.add_table(rows=len(pages_data), cols=4)
    tbl_pg.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pg.autofit = False
    
    for idx, (p_name, p_rt, p_elm, p_act) in enumerate(pages_data):
        row = tbl_pg.rows[idx]
        row.cells[0].width = Inches(1.3)
        row.cells[1].width = Inches(0.9)
        row.cells[2].width = Inches(2.2)
        row.cells[3].width = Inches(2.1)
        row.cells[0].paragraphs[0].text = p_name
        row.cells[1].paragraphs[0].text = p_rt
        row.cells[2].paragraphs[0].text = p_elm
        row.cells[3].paragraphs[0].text = p_act
        format_row(row, is_header=(idx == 0), is_alt=(idx % 2 == 1 and idx > 0))
        
    set_table_borders(tbl_pg, "CBD5E1")
    add_body_p(doc, "", space_after=4)

    # SECTION 28: MEERUT REGIONAL CONTEXT
    add_custom_heading(doc, "28. Meerut Regional Context & Authentic Demonstration Dataset", level=1)
    add_body_p(doc, "To ensure realistic testing during the Smart India Hackathon evaluation, ROADGUARD AI was calibrated with authentic geospatial data from the Meerut municipal zone (Uttar Pradesh). Meerut was selected because it represents a fast-growing National Capital Region (NCR) urban cluster featuring complex multi-agency road jurisdictions (UP PWD, Meerut Nagar Nigam, and NHAI):")
    
    add_bullet_p(doc, "Provincial Division Meerut (Executive Engineer: Er. R.K. Sharma). Responsible for major inter-district arterial corridors.", bold_prefix="1. Public Works Department (UP PWD Meerut): ")
    add_bullet_p(doc, "Municipal Corporation Limits (Chief Engineer: Er. A.K. Verma). Responsible for high-density intra-city residential and commercial streets.", bold_prefix="2. Meerut Nagar Nigam (Municipal Corporation): ")
    add_bullet_p(doc, "Project Implementation Unit Meerut (Project Director: Shri D.P. Gupta). Responsible for bypass expressways and national highway approaches.", bold_prefix="3. National Highways Authority of India (NHAI Meerut PIU): ")

    add_body_p(doc, "Eight Calibrated Meerut Road Segments:")
    add_bullet_p(doc, "1. Jail Chungi Road (MRT-RD-01 | Length: 3.8 km | Arterial | UP PWD | Health: 52 | Chronic Hotspot)")
    add_bullet_p(doc, "2. Surajkund Road (MRT-RD-02 | Length: 3.2 km | Major District | Nagar Nigam | Health: 61 | Chronic Hotspot)")
    add_bullet_p(doc, "3. Baghpat Bypass Road (MRT-RD-03 | Length: 7.5 km | Arterial | UP PWD | Health: 46 | Critical Deterioration)")
    add_bullet_p(doc, "4. Garh Road (MRT-RD-04 | Length: 5.6 km | Arterial | UP PWD | Health: 78 | Normal Condition)")
    add_bullet_p(doc, "5. Hapur Road (MRT-RD-05 | Length: 4.9 km | Arterial | Nagar Nigam | Health: 69 | Moderate Condition)")
    add_bullet_p(doc, "6. Kankerkheda Main Road (MRT-RD-06 | Length: 4.2 km | Major District | UP PWD | Health: 58 | Chronic Hotspot)")
    add_bullet_p(doc, "7. Civil Lines (MRT-RD-07 | Length: 3.1 km | Local Collector | Nagar Nigam | Health: 84 | Good Condition)")
    add_bullet_p(doc, "8. Malyana Chowki – Sharda Road Link (MRT-RD-08 | Length: 3.7 km | Local | Nagar Nigam | Health: 41 | Critical Deterioration)")

    add_body_p(doc, "Authentic Photographic Assets:")
    add_body_p(doc, "In accordance with strict hackathon ethics, all synthetic demo reports in the seed database use 6 real, reusable CC-licensed road distress photographs sourced from Wikimedia Commons and stored locally in `backend/uploads/demo/`:")
    add_bullet_p(doc, "`large-road-pothole.jpg` (Severe deep pothole hazard; CC BY-SA 4.0)")
    add_bullet_p(doc, "`driving-through-potholes.jpg` (Vehicle navigation through rutted asphalt; CC BY-SA 3.0)")
    add_bullet_p(doc, "`potholes-bengaluru-road.jpg` (Urban road surface crater; CC BY-SA 4.0)")
    add_bullet_p(doc, "`potholes-on-road.jpg` (Multiple clustered pavement depressions; CC BY-SA 3.0)")
    add_bullet_p(doc, "`waterlogged-pothole.jpg` (Monsoon ponding concealing subgrade pothole; CC BY-SA 4.0)")
    add_bullet_p(doc, "`repaired-road-patch.jpg` (Compacted asphalt maintenance patch; CC BY-SA 4.0)")

    add_callout(doc, "DEMONSTRATION DISCLAIMER: These Wikimedia Commons photographs are utilized strictly for prototype visual simulation and do not represent actual citizen-submitted photographs taken inside Meerut city.", title="DATA AUTHENTICITY DISCLAIMER", box_type="warning")

    doc.add_page_break()

print("Section 16-28 renderer ready.")
